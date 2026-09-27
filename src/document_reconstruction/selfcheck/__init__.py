"""Self-check: is this installation ready to convert documents?

    dre-selfcheck            (or: python -m document_reconstruction selfcheck)

Checks the OCR library and models, dictionaries, fonts, machine resources and speed,
then converts the control files next to this module and compares the result with
the expected text. Prints "OK" or the list of problems; exit code 0 or 1.
"""

from __future__ import annotations

import json
import re
import time
from collections import Counter
from pathlib import Path

HERE = Path(__file__).parent
MODELS = ("osd", "Latin", "Cyrillic")
WORD_LISTS = ("rus", "kaz", "kir", "uzb", "uzb_cyrl", "tgk", "aze", "eng")
HUNSPELL = ("ru_RU", "kk_KZ", "uz_UZ", "en_US")
FONTS = ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf")


def _words(text: str) -> Counter:
    text = re.sub("[‘’ʼ`']", "ʻ", text.lower().replace("ё", "е"))
    return Counter(re.findall(r"[^\W_]+(?:ʻ[^\W_]+)*", text))


def _docx_text(data: bytes) -> str:
    import io

    from docx import Document
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def run(say=print) -> list[str]:
    problems: list[str] = []

    def check(ok: bool, good: str, bad: str) -> None:
        say(("  yes  " if ok else "  NO   ") + (good if ok else bad))
        if not ok:
            problems.append(bad)

    say("Environment")
    try:
        from ..recognition.tesseract import TesseractBackend
        backend = TesseractBackend()
    except Exception as error:  # noqa: BLE001
        check(False, "", f"Tesseract does not start: {error}")
        return problems
    check(True, f"Tesseract {backend.runtime_version}, models: {backend.tessdata_dir}", "")
    missing = [model for model in MODELS if model not in backend.available]
    check(not missing, "OCR models: " + ", ".join(MODELS), "missing OCR models: " + ", ".join(missing))
    from ..recognition.lexicon import HUNSPELL_DIRECTORY, Lexicon, default_directory
    missing = [name for name in HUNSPELL if not (HUNSPELL_DIRECTORY / f"{name}.dic").exists()]
    check(not missing, "Hunspell dictionaries: " + ", ".join(HUNSPELL), "missing Hunspell dictionaries: " + ", ".join(missing))
    lexicon = Lexicon.ensure(backend.tessdata_dir)
    missing = [name for name in WORD_LISTS if name not in lexicon.languages]
    check(not missing, f"word lists ({default_directory()}): {len(lexicon.languages)}", "missing word lists: " + ", ".join(missing))
    from ..reconstruction.fit import FONT_DIRS
    missing = [font for font in FONTS if not any((Path(folder) / font).exists() for folder in FONT_DIRS)]
    check(not missing, "Liberation Serif fonts (Times New Roman metrics)", "missing fonts: " + ", ".join(missing))
    backend.close()

    from ..resources import cpu_count, memory_mb, page_threads
    memory = memory_mb()
    say(f"  cores: {cpu_count()}, memory: {memory if memory is not None else '?'} MB, pages at once: {page_threads()}")
    check(memory is None or memory >= 1800, "enough memory (1800 MB or more needed)", f"too little memory: {memory} MB (1800 MB or more needed)")

    say("Engine")
    from ..engine import ConversionEngine
    started = time.perf_counter()
    engine = ConversionEngine()
    engine.warm()
    say(f"  start {time.perf_counter() - started:.1f} s; speed: {engine.speed}")
    factor = float((engine.speed or {}).get("factor") or 1.0)
    forecast = engine.predictor.remaining_ms([_Plan(engine.predictor.page_ms(_ocr_route(), 9.0, True))]) / 1000
    check(forecast < 45, f"forecast for a one-page scan: {forecast:.1f} s (limit 45 s)",
          f"slow machine ({factor:.1f} times slower than the reference): forecast for a scan {forecast:.1f} s > 45 s, scans will be refused")

    say("Control files")
    expected = json.loads((HERE / "expected.json").read_text())
    for name, want in expected.items():
        started = time.perf_counter()
        result = engine.convert((HERE / name).read_bytes())
        seconds = time.perf_counter() - started
        reason = ((result.error or {}).get("details") or {}).get("reason") or (result.error or {}).get("code")
        if want["outcome"] == "reject":
            check(not result.ok and reason == want["reason"], f"{name}: refused ({reason}), {seconds:.1f} s",
                  f"{name}: expected refusal {want['reason']}, got: {'issued' if result.ok else reason}")
            continue
        if not result.ok:
            check(False, "", f"{name}: expected a DOCX, got refusal {reason} ({seconds:.1f} s)")
            continue
        text = _docx_text(result.docx)
        if "text" in want:
            reference, observed = _words("\n".join(want["text"])), _words(text)
            difference = sum(((reference - observed) + (observed - reference)).values())
            check(difference == 0, f"{name}: issued, text matches the reference, {seconds:.1f} s",
                  f"{name}: text differs from the reference by {difference} words")
        else:
            lost = [phrase for phrase in want["phrases"] if phrase not in text]
            check(not lost, f"{name}: issued, {seconds:.1f} s", f"{name}: missing phrases: {lost}")
        if seconds > 45:
            problems.append(f"{name}: {seconds:.1f} s, longer than 45 s")
    engine.close()
    return problems


class _Plan:
    def __init__(self, estimated_ms: float) -> None:
        from ..core import PageRoute
        self.route, self.estimated_ms, self.searchable_scan = PageRoute.OCR, estimated_ms, False


def _ocr_route():
    from ..core import PageRoute
    return PageRoute.OCR


def main(argv: list[str] | None = None) -> int:
    problems = run()
    print()
    if problems:
        print("PROBLEMS:")
        for problem in problems:
            print(" -", problem)
        return 1
    print("OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
