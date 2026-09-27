"""Native-table gaps must neither flatten genuine tables nor swallow annotations."""

import hashlib
import io
from pathlib import Path
import unittest
from unittest.mock import Mock

from docx import Document as WordDocument
from document_reconstruction.engine import ConversionEngine
from document_reconstruction.model import Box, Cell, Paragraph, Table
from document_reconstruction.recognition.types import RecognizedPage, Word
from document_reconstruction.understanding import understand_page


FIXTURE = Path(__file__).resolve().parents[1] / "benchmarks/stage23/native_table_gap_note.pdf"
SOURCE_SHA256 = "50487e49f2763ddff62b29fb1334318001e431427856c723cb7f5ba4c50b8779"


class NativeTableGapTests(unittest.TestCase):
    def test_table_gap_annotation_keeps_native_table_editable(self):
        source = FIXTURE.read_bytes()
        self.assertEqual(hashlib.sha256(source).hexdigest(), SOURCE_SHA256)
        ocr = Mock()
        result = ConversionEngine(ocr=ocr).convert(source)
        self.assertTrue(result.ok, result.error)
        ocr.recognize.assert_not_called()
        ocr.detect_script.assert_not_called()
        self.assertEqual(result.metrics["ocr_ms"], 0)
        tables = [block for block in result.model.pages[0].blocks if isinstance(block, Table)]
        self.assertEqual(len(tables), 1)
        self.assertEqual((tables[0].rows, tables[0].columns), (4, 3))
        paragraphs = [block.text for block in result.model.pages[0].blocks if isinstance(block, Paragraph)]
        self.assertIn("REF-7", paragraphs)
        self.assertEqual(result.model.text().count("REF-7"), 1)
        self.assertEqual(result.model.text().count("Approved"), 2)
        word = WordDocument(io.BytesIO(result.docx))
        self.assertEqual(len(word.tables), 1)
        self.assertIn("REF-7", [paragraph.text for paragraph in word.paragraphs])
        self.assertEqual(word.tables[0].cell(3, 2).text, "Approved")

    def test_unowned_gap_word_is_not_discarded_with_table_bounding_box(self):
        cells = [
            Cell("Header", 0, 0, box=Box(50, 100, 150, 112)),
            Cell("Value", 1, 0, box=Box(50, 142, 150, 154)),
        ]
        table = Table(Box(50, 100, 150, 154), 0, 2, 1, cells, source_region="native_borderless_table")
        words = [
            Word("Header", Box(60, 101, 110, 110), line_id=(0,), source_page=0),
            Word("REF-7", Box(60, 122, 100, 131), line_id=(1,), source_page=0),
            Word("Value", Box(60, 143, 110, 152), line_id=(2,), source_page=0),
        ]
        model_page = understand_page(RecognizedPage(0, 595, 842, words, tables=[table]))
        paragraphs = [block.text for block in model_page.blocks if isinstance(block, Paragraph)]
        self.assertIn("REF-7", paragraphs)
        self.assertEqual([block for block in model_page.blocks if isinstance(block, Table)][0].cells[1].text, "Value")
        self.assertEqual(" ".join([table.cells[0].text, *paragraphs, table.cells[1].text]).count("REF-7"), 1)


if __name__ == "__main__":
    unittest.main()
