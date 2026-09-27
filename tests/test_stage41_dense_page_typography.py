"""Keep ordinary Word typography legible while avoiding gratuitous spill pages."""
from __future__ import annotations

import io
import unittest

from docx import Document as WordDocument
from docx.shared import Pt

from document_reconstruction.model import Box, Cell, Document, Page, Paragraph, Table
from document_reconstruction.quality import verify
from document_reconstruction.reconstruction.docx import reconstruct


class DensePageTypographyTests(unittest.TestCase):
    @staticmethod
    def _document():
        paragraphs = [
            Paragraph(f"Editable source paragraph {index}", Box(45, 60 + index * 18, 500, 76 + index * 18), 0)
            for index in range(6)
        ]
        cells = [
            Cell(f"Editable cell {row}-{col}", row, col, box=Box(45+col*180, 200+row*25, 225+col*180, 225+row*25), source_page=0)
            for row in range(3) for col in range(2)
        ]
        table = Table(Box(45, 200, 405, 275), 0, 3, 2, cells)
        return Document([Page(0, 595, 842, paragraphs + [table])])

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_body_spacing_is_compact_without_shrinking_editable_text(self):
        model = self._document()
        payload = reconstruct(model)
        word = WordDocument(io.BytesIO(payload))
        # 6pt between every OCR-derived logical paragraph expands dense scans
        # into spill pages; 3pt remains readable and retains 11pt body type.
        self.assertEqual(word.styles['Body'].paragraph_format.space_after, Pt(3))
        self.assertEqual(word.styles['Body'].font.size, Pt(11))
        self.assertEqual(word.styles['Heading'].font.size, Pt(14))
        self.assertEqual(word.styles['Table Text'].font.size, Pt(10.5))
        self.assertEqual([p.text for p in word.paragraphs], [p.text for p in model.pages[0].blocks[:6]])
        self.assertEqual([[c.text for c in r.cells] for r in word.tables[0].rows],
                         [[f'Editable cell {row}-{col}' for col in range(2)] for row in range(3)])
        self.assertTrue(all(row.height is None for row in word.tables[0].rows))
        self.assertTrue(verify(payload, model).passed)

    def test_page_and_table_layout_are_not_replaced_by_fixed_dimensions(self):
        model = self._document()
        payload = reconstruct(model)
        word = WordDocument(io.BytesIO(payload))
        self.assertAlmostEqual(word.sections[0].page_width.pt, 595, delta=0.1)
        self.assertAlmostEqual(word.sections[0].page_height.pt, 842, delta=0.1)
        self.assertEqual(len(word.tables), 1)
        self.assertTrue(all(not p.paragraph_format.page_break_before for p in word.paragraphs))


if __name__ == '__main__':
    unittest.main()
