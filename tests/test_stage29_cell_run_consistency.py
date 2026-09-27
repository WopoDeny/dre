"""Source-backed table-cell run evidence must preserve exact editable text."""
import hashlib
import io
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
import unittest
from unittest.mock import Mock

from docx import Document as WordDocument
from document_reconstruction.engine import ConversionEngine
from document_reconstruction.model import Box, Cell, Document, Page, Table
from document_reconstruction.recognition.types import RecognizedPage, Word
from document_reconstruction.understanding import understand_page
from document_reconstruction.reconstruction.docx import reconstruct
from document_reconstruction.quality import verify, W_NS

PDF = Path(__file__).resolve().parents[1] / 'benchmarks/stage29/01_ru_kk_tender.pdf'
SOURCE_HASH = 'ebb93fdd2c5bcbbeede9d8349570a270664068b35a85e96467800a5ba035f0a4'

class Stage29CellRunTests(unittest.TestCase):
    def test_multiline_table_cell_does_not_attach_space_joined_runs(self):
        table = Table(Box(50, 50, 200, 120), 0, 1, 1, [
            Cell('72\n990,10', 0, 0, box=Box(60, 70, 175, 108), source_page=0),
        ], source_region='native_table')
        words = [
            Word('72', Box(68, 75, 89, 85), source_page=0, line_id=(0,)),
            Word('990,10', Box(68, 91, 113, 101), source_page=0, line_id=(1,)),
        ]
        page = understand_page(RecognizedPage(0, 595, 842, words, tables=[table]))
        tables = [b for b in page.blocks if isinstance(b, Table)]
        self.assertEqual(len(tables), 1)
        cell = tables[0].cells[0]
        self.assertEqual(cell.text, '72\n990,10')
        self.assertTrue(not cell.runs or ''.join(r.text for r in cell.runs) == cell.text)

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_editable_table_line_break_is_verified_not_silently_dropped(self):
        model = Document([Page(0, 595, 842, [
            Table(Box(50, 50, 300, 120), 0, 1, 1,
                  [Cell('72\n990,10', 0, 0)])
        ])])
        original = reconstruct(model)
        self.assertTrue(verify(original, model).passed)
        self.assertEqual(WordDocument(io.BytesIO(original)).tables[0].cell(0, 0).text, '72\n990,10')
        changed = io.BytesIO()
        removed = False
        with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(changed, 'w') as output:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == 'word/document.xml':
                    root = ET.fromstring(content)
                    for node in root.iter():
                        for child in list(node):
                            if child.tag == f'{{{W_NS}}}br':
                                node.remove(child)
                                removed = True
                    content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                output.writestr(entry, content)
        self.assertTrue(removed)
        failures = verify(changed.getvalue(), model).failures
        self.assertIn('table_cell_text_mismatch', failures)
        self.assertIn('unexpected_hard_line_breaks', failures)

    def test_frozen_source_preserves_multiline_native_cells_without_ocr(self):
        data = PDF.read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(), SOURCE_HASH)
        ocr = Mock()
        result = ConversionEngine(ocr=ocr).convert(data)
        self.assertTrue(result.ok, result.error)
        ocr.recognize.assert_not_called()
        ocr.detect_script.assert_not_called()
        self.assertEqual(len(result.model.integrity_issues()), 0)
        doc = WordDocument(io.BytesIO(result.docx))
        layout_tables = sum(1 for page in result.model.pages for block in page.blocks if isinstance(block, Table) and block.borderless)
        self.assertEqual(len(doc.tables) - layout_tables, 4)
        matching = [c.text for table in doc.tables for row in table.rows
                    for c in row.cells if 'C-20 |' in c.text]
        self.assertEqual(len(matching), 2)
        for value in matching:
            self.assertIn('72\n990,10', value)

if __name__ == '__main__':
    unittest.main()
