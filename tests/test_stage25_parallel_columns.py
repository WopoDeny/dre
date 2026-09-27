"""Source-backed, conservative Word layout for independent native columns."""

import hashlib
import io
import unittest
import zipfile
from copy import deepcopy
from pathlib import Path
from unittest.mock import Mock
from xml.etree import ElementTree as ET

from docx import Document as WordDocument

from document_reconstruction.engine import ConversionEngine
from document_reconstruction.quality import verify
from document_reconstruction.model import Box, Cell, Document, Page, Paragraph, Table, ColumnRelationship
from document_reconstruction.reconstruction.docx import reconstruct

ROOT = Path(__file__).resolve().parents[1]
PARALLEL = ROOT / "benchmarks/stage22/pdfs/14_native_four_script_columns.pdf"
GAP = ROOT / "benchmarks/stage23/native_table_gap_note.pdf"
PARALLEL_SHA = "f0fc805b55f5ce141601fdb6f6bc5c6d6f5e5a986a6ac7cf788bb6e5a7d17841"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def layout(docx: bytes):
    with zipfile.ZipFile(io.BytesIO(docx)) as data:
        return ET.fromstring(data.read("word/document.xml"))


def altered_xml_package(docx: bytes, change):
    """Change OOXML in a test without touching the source PDF or the model."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx)) as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "word/document.xml":
                xml = ET.fromstring(payload)
                change(xml)
                payload = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
            target.writestr(item, payload)
    return output.getvalue()


class NativeParallelColumnTests(unittest.TestCase):
    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')
    def test_frozen_four_script_columns_use_actual_word_columns_without_a_table(self):
        payload = PARALLEL.read_bytes()
        self.assertEqual(hashlib.sha256(payload).hexdigest(), PARALLEL_SHA)
        ocr = Mock()
        result = ConversionEngine(ocr=ocr).convert(payload)
        self.assertTrue(result.ok, result.error)
        ocr.recognize.assert_not_called()
        ocr.detect_script.assert_not_called()
        self.assertEqual(result.metrics["ocr_ms"], 0)
        self.assertEqual(result.metrics["retry_pages"], 0)
        self.assertEqual(len(result.model.pages[0].columns), 4)
        word = WordDocument(io.BytesIO(result.docx))
        self.assertEqual(len(word.tables), 0)
        self.assertEqual([p.text.replace("\n", "") for p in word.paragraphs if p.text.strip()],
                         [b.text for b in result.model.pages[0].blocks])
        xml = layout(result.docx)
        columns = xml.findall(".//w:sectPr/w:cols", NS)
        self.assertTrue(any(item.get(f"{{{NS['w']}}}num") == "4" for item in columns))
        self.assertEqual(len(xml.findall(".//w:br[@w:type='column']", NS)), 3)
        self.assertEqual(len(xml.findall(".//w:tbl", NS)), 0)
        self.assertEqual(len(xml.findall(".//w:drawing", NS)), 0)
        paragraph_elements = xml.findall(".//w:body/w:p", NS)
        column_starts = [p for p in paragraph_elements if p.find(".//w:br[@w:type='column']", NS) is not None]
        self.assertEqual(["".join(item.text or "" for item in p.findall(".//w:t", NS)) for p in column_starts],
                         ["Русский", "العربية", "中文"])
        self.assertTrue(result.quality["passed"])

    def test_table_and_unbalanced_column_metadata_do_not_trigger_word_columns(self):
        result = ConversionEngine(ocr=Mock()).convert(GAP.read_bytes())
        self.assertTrue(result.ok, result.error)
        self.assertEqual(len(WordDocument(io.BytesIO(result.docx)).tables), 1)
        self.assertFalse(any(item.get(f"{{{NS['w']}}}num") == "4"
                             for item in layout(result.docx).findall(".//w:sectPr/w:cols", NS)))

        page = Page(0, 595, 842, [
            Paragraph("Column A", Box(45, 100, 115, 112), 0, kind="heading", column=1),
            Paragraph("Some left-hand text", Box(45, 125, 170, 137), 0, column=1),
            Paragraph("Column B", Box(310, 100, 380, 112), 0, kind="heading", column=2),
            Paragraph("Some right-hand text", Box(310, 125, 460, 137), 0, column=2),
            Paragraph("Full-width explanatory note", Box(45, 180, 520, 194), 0, column=0),
        ], columns=[ColumnRelationship(1, Box(0, 100, 250, 137), 1),
                    ColumnRelationship(2, Box(250, 100, 595, 137), 2)])
        docx = reconstruct(Document([page]))
        self.assertEqual(len(WordDocument(io.BytesIO(docx)).tables), 0)
        self.assertEqual(len(layout(docx).findall(".//w:br[@w:type='column']", NS)), 0)
        self.assertFalse(any(item.get(f"{{{NS['w']}}}num") == "2"
                             for item in layout(docx).findall(".//w:sectPr/w:cols", NS)))

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_next_source_page_does_not_inherit_four_column_section(self):
        result = ConversionEngine(ocr=Mock()).convert(PARALLEL.read_bytes())
        self.assertTrue(result.ok, result.error)
        page = deepcopy(result.model.pages[0])
        page.index = 1
        page.columns = []
        page.blocks = [Paragraph("Next source page is single-column text.", Box(48, 90, 400, 105), 1)]
        combined = Document([result.model.pages[0], page])
        output = reconstruct(combined)
        report = verify(output, combined)
        self.assertTrue(report.passed, report.failures)
        self.assertEqual(report.logical_pages, 2)
        sections = layout(output).findall(".//w:sectPr/w:cols", NS)
        self.assertEqual([item.get(f"{{{NS['w']}}}num") for item in sections], ["1", "4", "1"])
        self.assertEqual(len(layout(output).findall(".//w:br[@w:type='column']", NS)), 3)

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_balanced_columns_without_spanning_heading_reuse_source_section(self):
        blocks = []
        relationships = []
        for column, x in enumerate((45, 310), start=1):
            relationships.append(ColumnRelationship(column, Box(x - 5, 100, x + 205, 171), column))
            for position, value in enumerate(("Heading", "First statement", "Second statement")):
                blocks.append(Paragraph(f"{value} {column}", Box(x, 100 + position * 24, x + 120, 112 + position * 24), 0,
                                        kind="heading" if position == 0 else "paragraph", column=column))
        model = Document([Page(0, 595, 842, blocks, columns=relationships)])
        output = reconstruct(model)
        report = verify(output, model)
        self.assertTrue(report.passed, report.failures)
        self.assertEqual(report.logical_pages, 1)
        columns = layout(output).findall(".//w:sectPr/w:cols", NS)
        self.assertEqual([item.get(f"{{{NS['w']}}}num") for item in columns], ["2"])
        self.assertEqual(len(layout(output).findall(".//w:br[@w:type='column']", NS)), 1)

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_quality_gate_rejects_broken_word_column_breaks_and_section_counts(self):
        result = ConversionEngine(ocr=Mock()).convert(PARALLEL.read_bytes())
        self.assertTrue(result.ok, result.error)
        original = result.docx
        model = result.model

        def break_to_line(xml):
            first = xml.find(".//w:br[@w:type='column']", NS)
            self.assertIsNotNone(first)
            first.set(f"{{{NS['w']}}}type", "textWrapping")

        broken_break = verify(altered_xml_package(original, break_to_line), model)
        self.assertFalse(broken_break.passed)
        self.assertIn("word_column_break_count_mismatch", broken_break.failures)
        self.assertIn("unexpected_hard_line_breaks", broken_break.failures)

        def four_to_one(xml):
            for item in xml.findall(".//w:sectPr/w:cols", NS):
                if item.get(f"{{{NS['w']}}}num") == "4":
                    item.set(f"{{{NS['w']}}}num", "1")
                    return
            self.fail("Missing source-backed four-column section")

        broken_section = verify(altered_xml_package(original, four_to_one), model)
        self.assertFalse(broken_section.passed)
        self.assertIn("word_column_section_mismatch", broken_section.failures)

        def relocate_break_without_changing_its_count(xml):
            paragraphs = xml.findall(".//w:body/w:p", NS)
            for index, paragraph in enumerate(paragraphs):
                text = "".join(element.text or "" for element in paragraph.findall(".//w:t", NS))
                if text == "Русский":
                    br = paragraph.find(".//w:br[@w:type='column']", NS)
                    self.assertIsNotNone(br)
                    for run in paragraph.findall("w:r", NS):
                        if br in list(run):
                            run.remove(br)
                            break
                    paragraphs[index - 1].findall("w:r", NS)[-1].append(br)
                    return
            self.fail("Column heading was not found")

        misplaced_break = verify(altered_xml_package(original, relocate_break_without_changing_its_count), model)
        self.assertFalse(misplaced_break.passed)
        self.assertIn("word_column_break_anchor_mismatch", misplaced_break.failures)


if __name__ == "__main__":
    unittest.main()
