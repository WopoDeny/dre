"""External ground truth must gate synthetic acceptance independently of DOCX creation."""
from __future__ import annotations
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from docx import Document as WordDocument
from scripts.stage29_evaluate import evaluate_case


class SyntheticGroundTruthAcceptanceTests(unittest.TestCase):
    def _case(self, *, add_table: bool):
        workspace=TemporaryDirectory()
        self.addCleanup(workspace.cleanup)
        root=Path(workspace.name)
        (root/'reference/ground_truth_pages').mkdir(parents=True)
        (root/'results').mkdir()
        truth={'tables_in_dom_order':[{'rows':[['ID','Value'],['A-07','٢٤']]}]}
        (root/'reference/ground_truth_pages/page1.json').write_text(json.dumps(truth))
        case={'file':'sample.pdf','page_modes':['native'],'pages':1,'ground_truth_pages':['ground_truth_pages/page1.json'],
              'exact_text_anchors':['Reference A-07']}
        doc=WordDocument();doc.add_paragraph('Reference A-07')
        if add_table:
            table=doc.add_table(rows=2,cols=2)
            table.cell(0,0).text='ID';table.cell(0,1).text='Value'
            table.cell(1,0).text='A-07';table.cell(1,1).text='٢٤'
        doc.save(root/'results/sample.docx')
        run={'id':1,'file':'sample.pdf','source_sha256':'frozen','pages':1,'page_modes':['native'],'ok':True,'error':None,
             'wall_ms':100,'anchor_count':1,'exact_anchor_hit_count':1,'missing_exact_anchors':[],
             'word_table_count':int(add_table),'docx_package_ok':True,'diagnostics':[{'ocr_performed':False}]}
        (root/'results/sample.json').write_text(json.dumps(run))
        return evaluate_case(case,root,root/'results')

    def test_missing_source_table_is_not_an_accepted_docx(self):
        result=self._case(add_table=False)
        self.assertTrue(result['conversion_ok'])
        self.assertEqual((result['table_expected'],result['table_actual']),(1,0))
        self.assertFalse(result['editability_accepted'])
        self.assertIn('table_structure_or_cell_content_mismatch',result['reason_not_accepted'])

    def test_matching_source_table_and_text_is_accepted(self):
        result=self._case(add_table=True)
        self.assertEqual(result['table_exact_matches'],1)
        self.assertTrue(result['editability_accepted'])

if __name__=='__main__':unittest.main()
