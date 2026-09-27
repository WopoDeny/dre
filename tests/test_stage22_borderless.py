"""Source-backed native borderless-table ownership regressions."""

import hashlib
import io
from pathlib import Path
import unittest
from unittest.mock import Mock

from docx import Document as WordDocument

from document_reconstruction.engine import ConversionEngine
from document_reconstruction.model import Paragraph, Table


FIXTURE = Path(__file__).resolve().parents[1] / "benchmarks/stage22/pdfs/18_native_borderless_table.pdf"
FIXTURE_SHA256 = "961194f05cb7720b78f2d46d35883292929ecdd361a0b6aecf73b631cc4ac90c"
PARALLEL_COLUMNS = FIXTURE.with_name("14_native_four_script_columns.pdf")
PARALLEL_SHA256 = "f0fc805b55f5ce141601fdb6f6bc5c6d6f5e5a986a6ac7cf788bb6e5a7d17841"


class NativeBorderlessOwnershipTests(unittest.TestCase):
    def test_source_backed_table_is_not_flattened_into_paragraphs(self):
        payload = FIXTURE.read_bytes()
        self.assertEqual(hashlib.sha256(payload).hexdigest(), FIXTURE_SHA256)
        ocr = Mock()
        result = ConversionEngine(ocr=ocr).convert(payload)
        self.assertTrue(result.ok, result.error)
        self.assertEqual(result.metrics["ocr_ms"], 0)
        ocr.recognize.assert_not_called()
        ocr.detect_script.assert_not_called()
        blocks = result.model.pages[0].blocks
        tables = [block for block in blocks if isinstance(block, Table)]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual((table.rows, table.columns), (4, 3))
        self.assertEqual(
            [[next(c.text for c in table.cells if c.row == row and c.column == col)
              for col in range(3)] for row in range(4)],
            [["Product", "Quantity", "Status"],
             ["Cable", "12", "Approved"],
             ["Adapter", "5", "Pending"],
             ["Monitor", "2", "Approved"]],
        )
        paragraphs = [block.text for block in blocks if isinstance(block, Paragraph)]
        self.assertIn("This note belongs below the table, outside its cells.", paragraphs)
        self.assertIn("Three repeated columns should become a table, not one long paragraph.", paragraphs)
        self.assertIn("Borderless inventory register", paragraphs)
        self.assertFalse(any("Cable" in text or "Quantity" in text for text in paragraphs))
        self.assertEqual(result.model.text().count("Approved"), 2)
        docx = WordDocument(io.BytesIO(result.docx))
        self.assertEqual(len(docx.tables), 1)
        self.assertEqual(len(docx.tables[0].rows), 4)
        self.assertEqual(docx.tables[0].cell(3, 2).text, "Approved")
        self.assertEqual(docx.tables[0].cell(0, 0).text, "Product")

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_four_independent_script_columns_are_not_promoted_to_a_table(self):
        """A parallel text layout is not a table despite repeated alignment."""
        payload = PARALLEL_COLUMNS.read_bytes()
        self.assertEqual(hashlib.sha256(payload).hexdigest(), PARALLEL_SHA256)
        ocr = Mock()
        result = ConversionEngine(ocr=ocr).convert(payload)
        self.assertTrue(result.ok, result.error)
        ocr.recognize.assert_not_called()
        self.assertFalse(any(isinstance(block, Table) for block in result.model.pages[0].blocks))
        self.assertEqual(len(result.model.pages[0].columns), 4)
        self.assertIn("已确认收货。", result.model.text())
        self.assertEqual(len(WordDocument(io.BytesIO(result.docx)).tables), 0)
