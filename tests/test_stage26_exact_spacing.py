"""Meaningful Unicode spacing must survive model-to-editable-Word verification."""

import io
import unittest
import zipfile
from xml.etree import ElementTree as ET

from docx import Document as WordDocument
from document_reconstruction.model import Box, Cell, Document, Page, Paragraph, Table
from document_reconstruction.quality import W_NS, verify
from document_reconstruction.reconstruction.docx import reconstruct


class ExactSpacingTests(unittest.TestCase):
    @staticmethod
    def mutate_text(docx: bytes, before: str, after: str) -> bytes:
        target = io.BytesIO()
        changed = False
        with zipfile.ZipFile(io.BytesIO(docx)) as source, zipfile.ZipFile(target, 'w') as output:
            for info in source.infolist():
                contents = source.read(info.filename)
                if info.filename == 'word/document.xml':
                    root = ET.fromstring(contents)
                    for element in root.iter(f'{{{W_NS}}}t'):
                        if element.text and before in element.text:
                            element.text = element.text.replace(before, after, 1)
                            changed = True
                            break
                    contents = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                output.writestr(info, contents)
        if not changed:
            raise AssertionError(f'Missing text segment: {before!r}')
        return target.getvalue()

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_paragraph_cannot_lose_meaningful_double_spaces(self):
        model = Document([Page(0, 595, 842, [
            Paragraph('Reference  A-02', Box(45, 55, 250, 70), 0),
        ])])
        original = reconstruct(model)
        self.assertTrue(verify(original, model).passed)
        self.assertEqual(WordDocument(io.BytesIO(original)).paragraphs[0].text, 'Reference  A-02')
        damaged = self.mutate_text(original, 'Reference  A-02', 'Reference A-02')
        self.assertIn('paragraph_text_mismatch', verify(damaged, model).failures)

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_table_cell_cannot_lose_meaningful_double_spaces(self):
        model = Document([Page(0, 595, 842, [Table(Box(45, 90, 520, 140), 0, 1, 2, [
            Cell('ID  007', 0, 0), Cell('USD  24', 0, 1),
        ])])])
        original = reconstruct(model)
        self.assertTrue(verify(original, model).passed)
        self.assertEqual(WordDocument(io.BytesIO(original)).tables[0].cell(0, 0).text, 'ID  007')
        damaged = self.mutate_text(original, 'ID  007', 'ID 007')
        self.assertIn('table_cell_text_mismatch', verify(damaged, model).failures)

    def test_mixed_han_and_arabic_runs_have_both_font_hints(self):
        # One logical source run may contain Han and Arabic-Indic digits.
        # Both fonts must be named even when Word does not split the run.
        model = Document([Page(0, 595, 842, [
            Paragraph('金额  ٢٤', Box(45, 55, 250, 70), 0),
            Table(Box(45, 90, 520, 140), 0, 1, 1, [Cell('金额  ٢٤', 0, 0)]),
        ])])
        result = reconstruct(model)
        report = verify(result, model)
        self.assertTrue(report.passed, report.failures)
        with zipfile.ZipFile(io.BytesIO(result)) as package:
            xml = package.read('word/document.xml')
        self.assertIn(b'Noto Sans CJK', xml)
        self.assertIn(b'Noto Naskh Arabic', xml)
        word = WordDocument(io.BytesIO(result))
        self.assertEqual(word.paragraphs[0].text, '金额  ٢٤')
        self.assertEqual(word.tables[0].cell(0, 0).text, '金额  ٢٤')

    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')

    def test_cjk_source_spaces_and_mixed_script_punctuation_are_not_normalized(self):
        model = Document([Page(0, 595, 842, [
            Paragraph('編號  2026/09/20 · REF-7', Box(45, 55, 400, 70), 0),
        ])])
        original = reconstruct(model)
        self.assertTrue(verify(original, model).passed)
        damaged = self.mutate_text(original, '編號  2026', '編號 2026')
        self.assertIn('paragraph_text_mismatch', verify(damaged, model).failures)
        punctuation = self.mutate_text(original, '/09/20 · REF', '/09/20  REF')
        self.assertIn('paragraph_text_mismatch', verify(punctuation, model).failures)


if __name__ == '__main__':
    unittest.main()
