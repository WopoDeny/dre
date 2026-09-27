"""Regression tests for script-aware Word East Asian font hints."""
from __future__ import annotations

import io
import unittest

from docx import Document as WordDocument
from docx.oxml.ns import qn

from document_reconstruction.model import Box, Cell, Document, InlineRun, Page, Paragraph, Table
from document_reconstruction.reconstruction.docx import reconstruct


class CjkVariantFontTests(unittest.TestCase):
    @staticmethod
    def paragraph_fonts(paragraph):
        return [(run.text, run._r.rPr.rFonts.get(qn("w:eastAsia")))
                for run in paragraph.runs if run.text and any("\u3400" <= c <= "\u9fff" for c in run.text)]

    def test_traditional_evidence_sets_coherent_font_for_common_and_variant_runs(self):
        box = Box(40, 40, 400, 65)
        model = Document([Page(0, 595, 842, [
            Paragraph("文件編號2026-A17，金額123元。", box, 0, runs=[
                InlineRun("文件", box, 0, script="Han"),
                InlineRun("編號", box, 0, script="Han"),
                InlineRun("2026-A17，", box, 0),
                InlineRun("金額123元。", box, 0, script="Han"),
            ]),
        ])])
        word = WordDocument(io.BytesIO(reconstruct(model)))
        self.assertEqual({font for _, font in self.paragraph_fonts(word.paragraphs[0])}, {"Noto Sans CJK TC"})
        self.assertEqual(word.paragraphs[0].text, model.pages[0].blocks[0].text)

    def test_simplified_text_remains_simplified_and_mixed_paragraph_is_not_force_converted(self):
        model = Document([Page(0, 595, 842, [
            Paragraph("文件编号，金额123元。", Box(40, 40, 400, 60), 0),
            Paragraph("文件編號 / 文件编号", Box(40, 80, 400, 100), 0),
        ])])
        word = WordDocument(io.BytesIO(reconstruct(model)))
        self.assertEqual({font for _, font in self.paragraph_fonts(word.paragraphs[0])}, {"Noto Sans CJK SC"})
        self.assertEqual([p.text for p in word.paragraphs], [p.text for p in model.pages[0].blocks])
        self.assertEqual(word.paragraphs[1].text, "文件編號 / 文件编号")

    def test_traditional_table_cell_inherits_variant_hint(self):
        model = Document([Page(0, 595, 842, [
            Table(Box(40, 40, 450, 120), 0, 1, 2,
                  [Cell("文件編號", 0, 0), Cell("编号金额", 0, 1)])
        ])])
        word = WordDocument(io.BytesIO(reconstruct(model)))
        first, second = word.tables[0].rows[0].cells
        self.assertEqual({font for _, font in self.paragraph_fonts(first.paragraphs[0])}, {"Noto Sans CJK TC"})
        self.assertEqual({font for _, font in self.paragraph_fonts(second.paragraphs[0])}, {"Noto Sans CJK SC"})

    def test_distinct_runs_on_one_mixed_variant_page_keep_independent_hints(self):
        box = Box(40, 40, 400, 65)
        model = Document([Page(0, 595, 842, [
            Paragraph("文件編號 / 文件编号", box, 0, runs=[
                InlineRun("文件編號", box, 0, script="Han"),
                InlineRun(" / ", box, 0),
                InlineRun("文件编号", box, 0, script="Han"),
            ])
        ])])
        paragraph = WordDocument(io.BytesIO(reconstruct(model))).paragraphs[0]
        self.assertEqual(self.paragraph_fonts(paragraph), [
            ("文件編號", "Noto Sans CJK TC"),
            ("文件编号", "Noto Sans CJK SC"),
        ])


if __name__ == "__main__":
    unittest.main()
