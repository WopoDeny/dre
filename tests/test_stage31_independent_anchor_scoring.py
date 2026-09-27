"""Frozen-source corpus scoring must inspect actual editable Word text and numbering."""
from __future__ import annotations
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from docx import Document
from docx.oxml.ns import qn
from scripts.stage29_evaluate import evaluate_case

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / 'benchmarks/stage29/stress20'
EVIDENCE = ROOT / 'release-evidence/runs/stage30/results'
CASE = json.loads((CORPUS / 'reference/manifest.json').read_text(encoding='utf-8'))['cases'][0]

class IndependentAnchorScoringTests(unittest.TestCase):
    def make_results(self):
        tmp = TemporaryDirectory(); self.addCleanup(tmp.cleanup)
        results = Path(tmp.name)
        stem = CASE['file'].removesuffix('.pdf')
        (results / (stem + '.json')).write_bytes((EVIDENCE / (stem + '.json')).read_bytes())
        doc = Document(EVIDENCE / (stem + '.docx'))
        return results, doc, results / (stem + '.docx')

    def test_list_numbers_verified_from_numbering_xml_not_missing_literal_text(self):
        results, doc, path = self.make_results()
        doc.save(path)
        assessment = evaluate_case(CASE, CORPUS, results)
        self.assertEqual(assessment['anchor_hits'], 19)
        self.assertEqual(len(assessment['numbering_verified_anchors']), 3)
        self.assertFalse(assessment['anchor_missing'])

    def test_damaged_visible_numbering_is_not_accepted(self):
        results, doc, path = self.make_results()
        for p in doc.paragraphs:
            if p.text.startswith('Подтвердить количество:') and p._p.xpath('./w:pPr/w:numPr'):
                num = p._p.xpath('./w:pPr/w:numPr')[0]
                num.getparent().remove(num)
        doc.save(path)
        assessment = evaluate_case(CASE, CORPUS, results)
        self.assertLess(assessment['anchor_hits'], assessment['anchor_total'])
        self.assertIn(CASE['exact_text_anchors'][3], assessment['anchor_missing'])
        self.assertFalse(assessment['editability_accepted'])

    def test_incorrect_marker_format_is_not_mistaken_for_source_number(self):
        results, doc, path = self.make_results()
        changed = False
        for abstract in doc.part.numbering_part.element.findall(qn('w:abstractNum')):
            for level in abstract.findall(qn('w:lvl')):
                marker = level.find(qn('w:lvlText'))
                if marker is not None and marker.get(qn('w:val')) == '%1.':
                    marker.set(qn('w:val'), '%1)')
                    changed = True
        self.assertTrue(changed)
        doc.save(path)
        assessment = evaluate_case(CASE, CORPUS, results)
        self.assertEqual(assessment['numbering_verified_anchors'], [])
        self.assertEqual(assessment['anchor_hits'], 16)
        self.assertFalse(assessment['editability_accepted'])

    def test_stale_json_does_not_override_damaged_docx_text(self):
        results, doc, path = self.make_results()
        anchor = CASE['exact_text_anchors'][0]
        changed = False
        for p in doc.paragraphs:
            if anchor in p.text:
                for run in p.runs:
                    if run.text:
                        run.text = 'MISSING SOURCE CONTENT'
                        changed = True
                        break
            if changed:
                break
        self.assertTrue(changed)
        doc.save(path)
        assessment = evaluate_case(CASE, CORPUS, results)
        self.assertIn(anchor, assessment['anchor_missing'])
        self.assertFalse(assessment['editability_accepted'])

if __name__ == '__main__': unittest.main()
