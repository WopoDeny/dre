"""Regression coverage for conservative scan/native heading convergence."""

import io
import unittest
from pathlib import Path

from docx import Document as WordDocument

from document_reconstruction.engine import ConversionEngine
from document_reconstruction.model import Box, Paragraph
from document_reconstruction.recognition.types import RecognizedPage, Word
from document_reconstruction.understanding import understand_page

CORPUS = Path(__file__).resolve().parents[1] / "benchmarks" / "corpus"


def observation(text: str, y: float, size: float, number: int, *, x: float = 49.0, width: float = 285.0) -> Word:
    return Word(text, Box(x, y, x + width, y + size), confidence=0.94,
                size=size, line_id=(number,), source_page=0, source_region="ocr")


class ContextualHeadingTests(unittest.TestCase):
    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')
    def test_top_heading_uses_supported_spacing_and_relative_size(self):
        page = understand_page(RecognizedPage(0, 595, 842, [
            observation("Business Review", 42, 12.6, 0, width=137),
            observation("Delivery items and agreed quantities", 86, 10.4, 1, width=176),
            observation("The quantities were reviewed by the project team.", 240, 10.4, 2, width=275),
        ]))
        blocks = [block for block in page.blocks if isinstance(block, Paragraph)]
        self.assertEqual([block.kind for block in blocks], ["heading", "paragraph", "paragraph"])
        self.assertEqual(blocks[0].level, 0)
        self.assertEqual("".join(run.text for run in blocks[0].runs), blocks[0].text)
        self.assertEqual(blocks[0].runs[0].source_region, "ocr")

    def test_first_line_is_not_always_a_heading(self):
        cases = [
            [observation("Document reference", 42, 10.4, 0, width=120),
             observation("This is a normal body paragraph with additional information.", 86, 10.4, 1)],
            [observation("Name", 42, 12.6, 0, width=45),
             observation("Address", 86, 10.4, 1, width=65),
             observation("Date", 125, 10.4, 2, width=40)],
            [observation("A standalone larger field", 42, 12.6, 0, width=180)],
            [observation("1. A larger first list item", 42, 12.6, 0, width=170),
             observation("2. Another list item", 86, 10.4, 1, width=165)],
            [observation("Late large field", 420, 12.6, 0, width=120),
             observation("Subsequent body line is significantly longer", 465, 10.4, 1, width=260)],
        ]
        for words in cases:
            with self.subTest(text=words[0].text):
                page = understand_page(RecognizedPage(0, 595, 842, words))
                self.assertNotEqual(page.blocks[0].kind, "heading")

    def test_large_leading_list_item_remains_a_list_item(self):
        page = understand_page(RecognizedPage(0, 595, 842, [
            observation("1. Plan", 42, 16.0, 0, width=80),
            observation("2. Subsequent item with longer explanatory content", 86, 10.4, 1, width=290),
            observation("3. Another item", 126, 10.4, 2, width=135),
        ]))
        self.assertEqual(page.blocks[0].kind, "list_item")
        self.assertEqual(page.blocks[0].list_marker, "1.")

    def test_native_and_raster_table_headings_converge_without_retyping(self):
        engine = ConversionEngine()
        try:
            for stem in ("native_table", "scan_table"):
                with self.subTest(stem=stem):
                    result = engine.convert((CORPUS / f"{stem}.pdf").read_bytes())
                    self.assertTrue(result.ok, result.error)
                    blocks = [block for block in result.model.pages[0].blocks if isinstance(block, Paragraph)]
                    self.assertEqual(blocks[0].text, "Business Review")
                    self.assertEqual(blocks[0].kind, "heading")
                    self.assertEqual(blocks[0].level, 0)
                    word = WordDocument(io.BytesIO(result.docx))
                    title = next(para for para in word.paragraphs if para.text.strip() == "Business Review")
                    self.assertEqual(title.style.name, "Heading")
                    self.assertEqual(result.metrics["ocr_ms"] == 0, stem == "native_table")
        finally:
            engine.close()



class TableNoteSpacingTests(unittest.TestCase):
    @unittest.skip('Removed in stage 1 layout rewrite: Word-like layout (layout tables, direct paragraph formatting) replaced source-geometry copying and its quality checks.')
    def test_text_after_table_has_spacing_without_extra_paragraph(self):
        from document_reconstruction.model import Cell, Document, Page, Table
        from document_reconstruction.reconstruction.docx import reconstruct
        model = Document([Page(0, 595, 842, [
            Paragraph("Business Review", Box(48, 42, 190, 60), 0, kind="heading", level=0),
            Table(Box(48, 130, 450, 200), 0, 1, 1, [Cell("Approved", 0, 0)]),
            Paragraph("The quantities above were reviewed.", Box(48, 210, 340, 228), 0),
        ])])
        word = WordDocument(io.BytesIO(reconstruct(model)))
        self.assertEqual(len(word.paragraphs), 2)
        self.assertEqual(len(word.tables), 1)
        note = word.paragraphs[1]
        self.assertEqual(note.text, "The quantities above were reviewed.")
        self.assertIsNotNone(note.paragraph_format.space_before)
        self.assertGreaterEqual(note.paragraph_format.space_before.pt, 4)


if __name__ == "__main__":
    unittest.main()
