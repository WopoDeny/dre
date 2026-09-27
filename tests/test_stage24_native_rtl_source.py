"""Source-backed rejection of demonstrably unreliable native RTL numeric extraction."""

import hashlib
import io
import json
import unittest
from pathlib import Path
from unittest.mock import Mock

import pymupdf
from docx import Document as WordDocument

from document_reconstruction.engine import ConversionEngine
from document_reconstruction.recognition.native import native_rtl_digit_order_anomalies

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "benchmarks" / "first_batch"
CASES = {case["id"]: case for case in json.loads((CORPUS / "manifest.json").read_text(encoding="utf-8"))["cases"]}


class NativeRTLSourceTests(unittest.TestCase):
    def test_frozen_native_arabic_fails_closed_instead_of_returning_corrupted_docx(self):
        case = CASES["10_native_arabic"]
        payload = (CORPUS / "pdfs" / case["filename"]).read_bytes()
        self.assertEqual(hashlib.sha256(payload).hexdigest(), case["sha256"])
        with pymupdf.open(stream=payload, filetype="pdf") as source:
            self.assertGreaterEqual(native_rtl_digit_order_anomalies(source[0]), 2)
        ocr = Mock()
        result = ConversionEngine(ocr=ocr).convert(payload)
        self.assertFalse(result.ok)
        self.assertIsNone(result.docx)
        self.assertEqual(result.error["code"], "UNREADABLE_DOCUMENT")
        self.assertEqual(result.error["details"]["reason"], "unreliable_native_rtl_digit_order")
        self.assertEqual(result.metrics["ocr_ms"], 0)
        self.assertEqual(result.metrics["retry_pages"], 0)
        ocr.recognize.assert_not_called()
        ocr.detect_script.assert_not_called()

    def test_proven_native_rtl_digit_order_is_not_rejected(self):
        # A correctly encoded Arabic numeric sequence advances in increasing X
        # even though the enclosing Arabic words advance right-to-left.
        text = {"blocks": [{"type": 0, "lines": [{"spans": [{"chars": [
            {"c": "ر", "bbox": (240, 20, 250, 38)},
            {"c": "ق", "bbox": (230, 20, 240, 38)},
            {"c": "م", "bbox": (220, 20, 230, 38)},
            {"c": " ", "bbox": (210, 20, 220, 38)},
            {"c": "٢", "bbox": (190, 20, 200, 38)},
            {"c": "٤", "bbox": (200, 20, 210, 38)},
        ]}]}]}]}
        self.assertEqual(native_rtl_digit_order_anomalies(text), 0)

    def test_latin_digits_do_not_trigger_arabic_indic_order_guard(self):
        text = {"blocks": [{"type": 0, "lines": [{"spans": [{"chars": [
            {"c": "ر", "bbox": (230, 20, 240, 38)},
            {"c": "ق", "bbox": (220, 20, 230, 38)},
            {"c": "م", "bbox": (210, 20, 220, 38)},
            {"c": " ", "bbox": (200, 20, 210, 38)},
            {"c": "2", "bbox": (190, 20, 200, 38)},
            {"c": "4", "bbox": (180, 20, 190, 38)},
        ]}]}]}]}
        self.assertEqual(native_rtl_digit_order_anomalies(text), 0)

    def test_native_kazakh_tajik_chinese_unchanged_and_never_ocr(self):
        for name in ["03_native_kazakh", "06_native_tajik", "13_native_chinese_traditional"]:
            with self.subTest(case=name):
                case = CASES[name]
                payload = (CORPUS / "pdfs" / case["filename"]).read_bytes()
                self.assertEqual(hashlib.sha256(payload).hexdigest(), case["sha256"])
                ocr = Mock()
                result = ConversionEngine(ocr=ocr).convert(payload)
                self.assertTrue(result.ok, result.error)
                self.assertIn(case["title"], WordDocument(io.BytesIO(result.docx)).paragraphs[0].text)
                self.assertEqual(result.metrics["ocr_ms"], 0)
                ocr.recognize.assert_not_called()

    def test_frozen_native_table_and_gap_and_four_columns_remain_editable(self):
        cases = [("benchmarks/stage22/pdfs/18_native_borderless_table.pdf", 1),
                 ("benchmarks/stage23/native_table_gap_note.pdf", 1),
                 ("benchmarks/stage22/pdfs/14_native_four_script_columns.pdf", 0)]
        for filename, table_count in cases:
            with self.subTest(filename=filename):
                ocr = Mock()
                result = ConversionEngine(ocr=ocr).convert((ROOT / filename).read_bytes())
                self.assertTrue(result.ok, result.error)
                data_tables = [t for t in WordDocument(io.BytesIO(result.docx)).tables if 'w:val="nil"' not in t._tbl.xml]
                self.assertEqual(len(data_tables), table_count)
                self.assertEqual(result.metrics["ocr_ms"], 0)
                ocr.recognize.assert_not_called()


if __name__ == "__main__":
    unittest.main()
